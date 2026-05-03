"""Build a gap-report Excel workbook from analyse_csv_gaps() output."""
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

_BLUE   = PatternFill("solid", fgColor="BDD7EE")  # pre-filled columns
_YELLOW = PatternFill("solid", fgColor="FFEB9C")  # user-fill columns
_INFO   = PatternFill("solid", fgColor="DDEEFF")  # instruction rows
_SEC_CLIENT  = PatternFill("solid", fgColor="E2EFDA")
_SEC_PROJECT = PatternFill("solid", fgColor="FFF2CC")


def _cell(ws, row, col, value, fill=None, bold=False, wrap=True, center=False):
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    c.font = Font(bold=bold)
    align_kw = {"wrap_text": wrap, "vertical": "top"}
    if center:
        align_kw["horizontal"] = "center"
    c.alignment = Alignment(**align_kw)
    return c


def build_gap_excel(gaps: dict, row_counts: dict) -> bytes:
    """
    gaps       — output of db.analyse_csv_gaps()
    row_counts — {(client_code, client_suffix): n_rows}
    Returns Excel file content as bytes.
    """
    wb = Workbook()

    # ------------------------------------------------------------------ #
    # Sheet A — client exists, project code missing                        #
    # ------------------------------------------------------------------ #
    ws_a = wb.active
    ws_a.title = "A - Add Project Code"

    A_PRE  = ["client_code", "client_name", "client_suffix",
               "rows_in_csv", "existing_projects (read-only)"]
    A_USER = ["project_to_attach_to", "code_name", "code_description",
               "budget_amount", "status", "date_start (YYYY-MM-DD)", "date_end (YYYY-MM-DD)"]
    A_COLS = A_PRE + A_USER

    ws_a.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(A_COLS))
    _cell(ws_a, 1, 1,
          "Group A — Client already exists. Add the missing Project Code on page 6. "
          "Blue = pre-filled  |  Yellow = you fill in  |  "
          "'project_to_attach_to' must match an existing project name exactly.",
          fill=_INFO, bold=True)
    ws_a.row_dimensions[1].height = 36

    for ci, col in enumerate(A_COLS, 1):
        _cell(ws_a, 2, ci, col,
              fill=_BLUE if ci <= len(A_PRE) else _YELLOW,
              bold=True, center=True)

    items_a = sorted(gaps["missing_code"], key=lambda x: (x["client_code"], x["client_suffix"]))
    if items_a:
        for ri, item in enumerate(items_a, 3):
            cc, cs = item["client_code"], item["client_suffix"]
            pre  = [cc, item["client_name"], cs,
                    row_counts.get((cc, cs), 0),
                    " | ".join(item["existing_projects"])]
            user = ["", "", "", 0, "Active", "", ""]
            for ci, v in enumerate(pre + user, 1):
                _cell(ws_a, ri, ci, v, fill=_BLUE if ci <= len(A_PRE) else _YELLOW)
    else:
        _cell(ws_a, 3, 1, "(none — all clients with matching codes are already set up)")

    for i, w in enumerate([14, 20, 12, 10, 45, 30, 22, 28, 14, 20, 22, 22], 1):
        ws_a.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------ #
    # Sheet B — client missing entirely                                    #
    # ------------------------------------------------------------------ #
    ws_b = wb.create_sheet("B - Add Client + Code")

    B_PRE     = ["client_code", "client_suffix", "rows_in_csv", "is_internal (0009xxx)"]
    B_CLIENT  = ["client_name", "client_name_for_invoices", "vat_number"]
    B_PROJECT = ["project_name", "project_description", "project_status"]
    B_CODE    = ["code_name", "code_description", "budget_amount",
                 "code_status", "date_start (YYYY-MM-DD)", "date_end (YYYY-MM-DD)"]
    B_COLS    = B_PRE + B_CLIENT + B_PROJECT + B_CODE

    n_pre = len(B_PRE)
    n_cli = len(B_CLIENT)
    n_prj = len(B_PROJECT)

    ws_b.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(B_COLS))
    _cell(ws_b, 1, 1,
          "Group B — Client not in DB. Create Client (page 3 or 11) → Project → Project Code. "
          "Blue = pre-filled  |  Yellow = you fill in  |  "
          "Internal rows (0009xxx) are non-billable overhead — add only if you want to track them.",
          fill=_INFO, bold=True)
    ws_b.row_dimensions[1].height = 36

    # section label row
    sections = [
        (1,                         n_pre,               "PRE-FILLED",    _BLUE),
        (n_pre + 1,                 n_pre + n_cli,       "CLIENT",        _SEC_CLIENT),
        (n_pre + n_cli + 1,         n_pre + n_cli + n_prj, "PROJECT",     _SEC_PROJECT),
        (n_pre + n_cli + n_prj + 1, len(B_COLS),         "PROJECT CODE",  _YELLOW),
    ]
    for sc, ec, label, sfill in sections:
        ws_b.merge_cells(start_row=2, start_column=sc, end_row=2, end_column=ec)
        _cell(ws_b, 2, sc, label, fill=sfill, bold=True, center=True)

    for ci, col in enumerate(B_COLS, 1):
        _cell(ws_b, 3, ci, col,
              fill=_BLUE if ci <= n_pre else _YELLOW,
              bold=True, center=True)

    items_b = sorted(gaps["missing_client"],
                     key=lambda x: (not x["is_internal"], x["client_code"], x["client_suffix"]))
    if items_b:
        for ri, item in enumerate(items_b, 4):
            cc, cs = item["client_code"], item["client_suffix"]
            pre  = [cc, cs, row_counts.get((cc, cs), 0),
                    "Yes" if item["is_internal"] else "No"]
            user = ["", "", "",       # client
                    "", "", "Active", # project
                    "", "", 0, "Active", "", ""]  # code
            for ci, v in enumerate(pre + user, 1):
                _cell(ws_b, ri, ci, v, fill=_BLUE if ci <= n_pre else _YELLOW)
    else:
        _cell(ws_b, 4, 1, "(none — all clients are already in the database)")

    col_widths_b = [14, 12, 10, 16, 24, 24, 16, 24, 24, 18, 20, 24, 14, 18, 22, 22]
    for i, w in enumerate(col_widths_b, 1):
        ws_b.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
